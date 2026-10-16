"""Builds the partner pipeline workbook from the researched CSVs."""
import csv, glob, re, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter

D = os.path.dirname(os.path.abspath(__file__))
OUT = sys.argv[1] if len(sys.argv) > 1 else f"{D}/pipeline.xlsx"

WAVE1 = ["Ossington", "Queen West", "King West", "Financial District", "Yorkville", "Leslieville"]
ORDER = WAVE1 + ["Little Italy", "Kensington and Chinatown", "The Annex", "Roncesvalles", "Danforth", "Liberty Village",
                 "Little Portugal", "Parkdale", "The Junction", "St. Clair West",
                 "Distillery and Corktown", "Summerhill and Rosedale", "Midtown", "Leaside", "Riverdale"]
CATS = ["Restaurant", "Brunch", "Café", "Bakery", "Spa", "Studio", "Hotel restaurant or spa", "Experience"]
STATUS = ["Target", "Contacted", "Meeting booked", "Met", "LOI signed", "Agreement signed", "Not now"]
HOW = ["Walked in", "Email", "Instagram", "Friend", "Referral", "Event"]
OPEN = ["Not checked", "Confirmed open", "Closed"]

# Rows the researchers flagged as weak evidence of being open, or worth a note.
WEAK = """Bindia|Volos|Little Anthony's|Edulis|Gusto 101|Le Select|Eastside Social|Joso's|Kasa Moto|Sofia|The Oxley|Amal|Fieramosca|Museum Tavern|KOKO! Share Bar|Lady Marmalade|Café Boulud|Ricky + Olivia|Hume|Aroi Thai|Lamoon|Yakitori Momiji|Cafe Rosarium|Mom's Recipes|DaiLo|Graination|Qin's Garden|Rikki Tikki|Bunner's|Le Paradis|Fat Pasha|Chadwick's|Bar Mercurio|The Stop|Fantail|Mabel's|Dil Se|Sorry I've Got Plants|Cherry Bomb|Athens Pastries|Brasserie Côte|The Haifa Room|Azhar|Manita|Bar Prima|Aleria|The Lakeview|MSSM Ossington|Côte de Boeuf|Fat Choi|Papi Chulo's|Papi & Rosita|The Good Son|Trinity Market|Little Ese|Archeo|Aafrien|Namul|White Lily Diner|Bonjour Brioche|Lake Inez|El Cazo|Agak Agak|Pii Nong Thai|Sukhothai|Botham's|Nodo|Nuna|Afrobeat Kitchen|Easy|Batibot|Himalayan Kitchen|Loga's Corner|Tiny Cafe|Shelíss|Daymi's|Octopus Garden Yoga|Living Water Therapies|Coco Reformer Pilates|Galeria Pilates + Coffee|Stillwater Spa""".split("|")
norm = lambda s: re.sub(r"[^a-z0-9]", "", s.lower().replace("&", "and"))
WEAKN = {norm(w) for w in WEAK}

rows, seen = [], set()
for f in sorted(glob.glob(f"{D}/*.csv")):
    with open(f, encoding="utf-8") as fh:
        for r in csv.DictReader(fh):
            r = {k.strip(): (v or "").strip() for k, v in r.items() if k}
            if not r.get("name"):
                continue
            key = norm(r["name"]) + "|" + norm(r.get("address", ""))[:6]
            if norm(r["name"]) in {norm(x["name"]) for x in rows} or key in seen:
                continue
            seen.add(key)
            for k, v in list(r.items()):
                r[k] = v.replace("—", ",").replace("–", "-")
            rows.append(r)

nb_rank = {n: i for i, n in enumerate(ORDER)}
cat_rank = {c: i for i, c in enumerate(CATS)}
rows.sort(key=lambda r: (nb_rank.get(r["neighbourhood"], 99), cat_rank.get(r.get("category", ""), 50), r["name"].lower()))

wb = Workbook()
ink, copper, mute = "16140F", "9A5A32", "6E685E"
hfill = PatternFill("solid", fgColor="1B1A17")
thin = Side(style="thin", color="E6E1D9")

# ---- Pipeline sheet
ws = wb.active
ws.title = "Pipeline"
cols = [("#", 5), ("Wave", 7), ("Place", 28), ("Neighbourhood", 20), ("Category", 14), ("Cuisine or type", 20),
        ("Address", 30), ("Website", 28), ("Instagram", 18), ("Phone", 15), ("Price", 7), ("Why it fits", 48),
        ("Open?", 15), ("Status", 16), ("Contact name", 18), ("Role", 12), ("Contact email or phone", 24),
        ("How we know them", 15), ("Last contact", 12), ("Next step", 26), ("Next step date", 13),
        ("Slow nights", 14), ("Offer idea", 26), ("Notes", 34), ("Source", 30)]
C = {name: i + 1 for i, (name, _) in enumerate(cols)}
for i, (name, w) in enumerate(cols, 1):
    c = ws.cell(1, i, name)
    c.font = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
    c.fill = hfill
    c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.column_dimensions[get_column_letter(i)].width = w
ws.row_dimensions[1].height = 30

def link(cell, url):
    if url and url.startswith("http"):
        cell.hyperlink = url
        cell.font = Font(color="1F5FAF", underline="single")

for n, r in enumerate(rows, 1):
    rr = n + 1
    nb = r["neighbourhood"]
    vals = {"#": n, "Wave": 1 if nb in WAVE1 else 2, "Place": r["name"], "Neighbourhood": nb,
            "Category": r.get("category", ""), "Cuisine or type": r.get("cuisine", ""), "Address": r.get("address", ""),
            "Website": r.get("website", ""), "Instagram": r.get("instagram", ""), "Phone": r.get("phone", ""),
            "Price": r.get("price", ""), "Why it fits": r.get("fit_note", ""), "Open?": "Not checked",
            "Status": "Target", "Notes": "Check it's open before visiting: weak online evidence." if any(norm(r["name"]).startswith(w) for w in WEAKN if len(w) > 3) else "",
            "Source": r.get("source_url", "")}
    for k, v in vals.items():
        c = ws.cell(rr, C[k], v)
        c.alignment = Alignment(vertical="top", wrap_text=k in ("Why it fits", "Notes", "Next step", "Offer idea"))
        c.border = Border(bottom=thin)
    ws.cell(rr, C["Place"]).font = Font(bold=True)
    link(ws.cell(rr, C["Website"]), r.get("website", ""))
    link(ws.cell(rr, C["Source"]), r.get("source_url", ""))
    ig = r.get("instagram", "").lstrip("@")
    if ig and not ig.startswith("http"):
        ws.cell(rr, C["Instagram"]).value = "@" + ig.split("/")[-1].strip("/") if "/" in ig else "@" + ig
        link(ws.cell(rr, C["Instagram"]), "https://instagram.com/" + ig.split("/")[-1].strip("/"))
    elif ig:
        handle = ig.rstrip("/").split("/")[-1]
        ws.cell(rr, C["Instagram"]).value = "@" + handle
        link(ws.cell(rr, C["Instagram"]), ig)
    for k in ("Last contact", "Next step date"):
        ws.cell(rr, C[k]).number_format = "mmm d"

last = len(rows) + 1 + 200  # room to add places
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(rows) + 1}"

# ---- Lists sheet
L = wb.create_sheet("Lists")
for j, (title, items) in enumerate([("Status", STATUS), ("Neighbourhood", ORDER), ("Category", CATS), ("How we know them", HOW), ("Open?", OPEN)], 1):
    L.cell(1, j, title).font = Font(bold=True)
    for i, v in enumerate(items, 2):
        L.cell(i, j, v)
    L.column_dimensions[get_column_letter(j)].width = 24

def dv(col, listcol, n):
    v = DataValidation(type="list", formula1=f"=Lists!${listcol}$2:${listcol}${n + 1}", allow_blank=True,
                       showErrorMessage=True, errorTitle="Not in the list", error="Pick one from the list.")
    L_ = get_column_letter(C[col])
    v.add(f"{L_}2:{L_}{last}")
    ws.add_data_validation(v)
dv("Status", "A", len(STATUS)); dv("Neighbourhood", "B", len(ORDER)); dv("Category", "C", len(CATS))
dv("How we know them", "D", len(HOW)); dv("Open?", "E", len(OPEN))
dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True)
for k in ("Last contact", "Next step date"):
    dvd.add(f"{get_column_letter(C[k])}2:{get_column_letter(C[k])}{last}")
ws.add_data_validation(dvd)

S = get_column_letter(C["Status"])
colours = {"Contacted": "FFF3D6", "Meeting booked": "FDE7CF", "Met": "E8E0F5", "LOI signed": "DDF0E4",
           "Agreement signed": "BFE5CC", "Not now": "EEEEEE"}
for s, col in colours.items():
    ws.conditional_formatting.add(f"A2:{get_column_letter(len(cols))}{last}",
        FormulaRule(formula=[f'${S}2="{s}"'], fill=PatternFill("solid", fgColor=col)))
ND = get_column_letter(C["Next step date"])
ws.conditional_formatting.add(f"{ND}2:{ND}{last}",
    FormulaRule(formula=[f'AND({ND}2<>"",{ND}2<TODAY(),${S}2<>"Agreement signed",${S}2<>"Not now")'],
                font=Font(color="C0392B", bold=True)))
OP = get_column_letter(C["Open?"])
ws.conditional_formatting.add(f"{OP}2:{OP}{last}", CellIsRule(operator="equal", formula=['"Closed"'], font=Font(color="999999", strike=True)))

# ---- Start here sheet
H = wb.create_sheet("Start here", 0)
H.column_dimensions["A"].width = 30
for c in "BCDEFGHI":
    H.column_dimensions[c].width = 13
H["A1"] = "The 6 Pass · partner pipeline"
H["A1"].font = Font(size=18, bold=True, color=ink)
H["A2"] = f"{len(rows)} places to approach, researched October 2026 from public sources. Counts below update by themselves."
H["A2"].font = Font(color=mute)
lines = [
    "How to use it",
    "1. Work one neighbourhood at a time. Wave 1 is the six launch neighbourhoods; start with Ossington.",
    "2. Before walking in, check the place is still open (Open? column). Some rows note weak online evidence.",
    "3. Every row you touch gets a Status, a Next step and a Next step date. Overdue dates turn red.",
    "4. Write their slow nights and offer idea as you hear them. That becomes their offer in the app.",
    "5. Be picky: we need 50 Founding Partners, not 300. A 'Not now' with a reason is a good outcome.",
    "Contacts are blank on purpose: ask for the owner or general manager in person, then fill them in.",
]
for i, t in enumerate(lines, 4):
    H.cell(i, 1, t).font = Font(bold=(i == 4), color=ink if i == 4 else "333333")

r0 = 13
H.cell(r0, 1, "By status").font = Font(bold=True, color=copper)
for i, s in enumerate(STATUS, r0 + 1):
    H.cell(i, 1, s)
    H.cell(i, 2, f'=COUNTIF(Pipeline!{S}:{S},"{s}")')
H.cell(r0 + 1 + len(STATUS), 1, "Goal for the first week: 10 conversations, 5 demos, 2 letters of intent.").font = Font(italic=True, color=mute)

r1 = r0 + len(STATUS) + 3
NB = get_column_letter(C["Neighbourhood"]); CA = get_column_letter(C["Category"])
hdr = ["Neighbourhood", "Places", "Contacted+", "Met", "LOI signed", "Signed"]
for j, t in enumerate(hdr, 1):
    c = H.cell(r1, j, t); c.font = Font(bold=True, color="FFFFFF"); c.fill = hfill
present = [n for n in ORDER if any(r["neighbourhood"] == n for r in rows)]
for i, n in enumerate(present, r1 + 1):
    H.cell(i, 1, n + ("  (wave 1)" if n in WAVE1 else ""))
    rng = f"Pipeline!{NB}:{NB}"
    H.cell(i, 2, f'=COUNTIF({rng},"{n}")')
    H.cell(i, 3, f'=COUNTIFS({rng},"{n}",Pipeline!{S}:{S},"<>Target",Pipeline!{S}:{S},"<>")')
    H.cell(i, 4, f'=COUNTIFS({rng},"{n}",Pipeline!{S}:{S},"Met")')
    H.cell(i, 5, f'=COUNTIFS({rng},"{n}",Pipeline!{S}:{S},"LOI signed")')
    H.cell(i, 6, f'=COUNTIFS({rng},"{n}",Pipeline!{S}:{S},"Agreement signed")')
    if n in WAVE1:
        H.cell(i, 1).font = Font(bold=True)
tot = r1 + 1 + len(present)
H.cell(tot, 1, "Total").font = Font(bold=True)
for j in range(2, 7):
    col = get_column_letter(j)
    H.cell(tot, j, f"=SUM({col}{r1 + 1}:{col}{tot - 1})").font = Font(bold=True)

r2 = tot + 2
H.cell(r2, 1, "By category").font = Font(bold=True, color=copper)
for i, cat in enumerate([c for c in CATS if any(r.get("category") == c for r in rows)], r2 + 1):
    H.cell(i, 1, cat); H.cell(i, 2, f'=COUNTIF(Pipeline!{CA}:{CA},"{cat}")')

wb.move_sheet("Lists", offset=1)
wb.save(OUT)
from collections import Counter
print(len(rows), "rows ->", OUT)
print(Counter(r["neighbourhood"] for r in rows))
print(Counter(r.get("category") for r in rows))
